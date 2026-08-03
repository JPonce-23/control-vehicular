from docxtpl import DocxTemplate
from docx import Document
from datetime import datetime
import os


def check(valor_actual, valor_esperado):
    if valor_actual is None:
        return ""
    return "✓" if str(valor_actual).lower().strip() == valor_esperado else ""


def normalizar_gasolina(valor_actual):
    if valor_actual is None:
        return ""
    
    valor_actual = str(valor_actual).lower().strip()

    equivalencias = {
        "vacio": "0",
        "vacío": "0",
        "0": "0",
        "cuarto": "1/4",
        "1/4": "1/4",
        "medio": "1/2",
        "1/2": "1/2",
        "tres cuartos": "3/4",
        "tres_cuartos": "3/4",
        "3/4": "3/4",
        "lleno": "4/4",
        "4/4": "4/4"
    }

    return equivalencias.get(valor_actual, valor_actual)


def normalizar_llantas(valor_actual):
    return normalizar_gasolina(valor_actual)


def formatear_fecha(fecha):
    if fecha is None:
        return "N/A"
    return fecha.strftime("%d/%m/%Y")


def mes_en_letra(fecha):
    if fecha is None:
        return ""

    meses = {
        1: "enero",
        2: "febrero",
        3: "marzo",
        4: "abril",
        5: "mayo",
        6: "junio",
        7: "julio",
        8: "agosto",
        9: "septiembre",
        10: "octubre",
        11: "noviembre",
        12: "diciembre"
    }
        
    return meses[fecha.month]


def valor(campo, default=""):
    if campo is None:
        return default

    texto = str(campo).strip()

    if texto.lower() == "string":
        return default

    return texto


def normalizar_clave(texto):
    texto = str(texto).lower().strip()
    reemplazos = {
        "á": "a",
        "é": "e",
        "í": "i",
        "ó": "o",
        "ú": "u",
        "ñ": "n"
    }

    for original, reemplazo in reemplazos.items():
        texto = texto.replace(original, reemplazo)

    texto = texto.replace(" ", "_")
    return texto


def generar_checks_condiciones(condiciones):
    checks = {}

    for revision, item in condiciones:
        clave = normalizar_clave(item.nombre)
        estado = str(revision.estado).lower().strip()

        checks[f"{clave}_b"] = "✓" if estado == "bueno" else ""
        checks[f"{clave}_r"] = "✓" if estado == "regular" else ""
        checks[f"{clave}_m"] = "✓" if estado == "malo" else ""

    return checks


def generar_checks_inventario(inventario):
    checks = {}

    for revision, item in inventario:
        clave = normalizar_clave(item.nombre)
        estado = str(revision.estado).lower().strip()

        if estado == "correcto":
            checks[clave] = "✓"
        elif estado == "na":
            checks[clave] = "N/A"
        else:
            checks[clave] = ""

    return checks



def _generar_resguardo_basico(ruta_archivo, contexto, condiciones, inventario):
    """Genera un documento funcional cuando no existe la plantilla institucional."""
    documento = Document()
    documento.add_heading("Resguardo de vehículo oficial", level=1)

    tabla = documento.add_table(rows=0, cols=2)
    tabla.style = "Table Grid"
    campos = [
        ("Fecha de salida", contexto.get("fecha_inicio")),
        ("Fecha estimada de regreso", contexto.get("fecha_fin")),
        ("Persona autorizada", contexto.get("nombre_persona")),
        ("Cargo", contexto.get("cargo")),
        ("Área", contexto.get("area_adscripcion")),
        ("Vehículo", f"{contexto.get('marca', '')} {contexto.get('tipo', '')}".strip()),
        ("Modelo", contexto.get("modelo")),
        ("Placa", contexto.get("placa")),
        ("Número de serie", contexto.get("num_serie")),
        ("Número económico", contexto.get("num_economico")),
        ("Tarjeta de gasolina", contexto.get("num_tarjeta_gasolina")),
        ("Finalidad", contexto.get("finalidad_uso")),
        ("Kilometraje de salida", contexto.get("km_odometro_salida")),
        ("Nivel de gasolina", contexto.get("nivel_gasolina")),
        ("Nivel de llantas", contexto.get("estado_llantas")),
    ]
    for etiqueta, dato in campos:
        celdas = tabla.add_row().cells
        celdas[0].text = str(etiqueta)
        celdas[1].text = str(dato or "")

    if condiciones:
        documento.add_heading("Condiciones", level=2)
        for revision, item in condiciones:
            documento.add_paragraph(
                f"{item.nombre}: {revision.estado}"
                + (f" — {revision.observaciones}" if revision.observaciones else "")
            )

    if inventario:
        documento.add_heading("Inventario", level=2)
        for revision, item in inventario:
            documento.add_paragraph(
                f"{item.nombre}: {revision.estado}"
                + (f" — {revision.observaciones}" if revision.observaciones else "")
            )

    documento.add_paragraph("Nombre y firma de quien recibe: ______________________________")
    documento.add_paragraph("Nombre y firma de quien entrega: _____________________________")
    documento.save(ruta_archivo)

def generar_resguardo_word(salida, vehiculo, persona, condiciones=None, inventario=None):
    BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))

    carpeta_resguardos = os.path.join(BASE_DIR, "archivos", "resguardos")
    ruta_template = os.path.join(BASE_DIR, "templates", "resguardo_template.docx")

    os.makedirs(carpeta_resguardos, exist_ok=True)

    nombre_archivo = f"resguardo_salida_{salida.id}.docx"
    ruta_archivo = os.path.join(carpeta_resguardos, nombre_archivo)

    nombre_persona = (
        f"{valor(persona.nombre, '')} "
        f"{valor(persona.apellido_paterno, '')} "
        f"{valor(persona.apellido_materno, '')}"
    ).strip()

    contexto = {
        "dia": datetime.now().strftime("%d"),
        "mes": datetime.now().strftime("%m"),
        "anio": datetime.now().strftime("%Y"),

        "fecha_salida": salida.fecha_salida.day if salida.fecha_salida else "",
        "fecha_regreso": salida.fecha_regreso_estimada.day if salida.fecha_regreso_estimada else "",
        "mes_regreso": mes_en_letra(salida.fecha_regreso_estimada) if salida.fecha_regreso_estimada else "",
        "anio_regreso": salida.fecha_regreso_estimada.year if salida.fecha_regreso_estimada else "",

        "marca": valor(vehiculo.marca),
        "tipo": valor(vehiculo.tipo),
        "modelo": valor(vehiculo.modelo_anio),
        "cilindros": valor(vehiculo.cilindros),
        "num_motor": valor(vehiculo.num_motor),
        "num_serie": valor(vehiculo.num_serie),
        "num_poliza": valor(vehiculo.num_poliza),
        "placa": valor(vehiculo.placa),
        "num_inventario": "",
        "num_economico": "",
        "color": valor(vehiculo.color),

        "tipo_movimiento": "",
        "forma_movimiento": "PROVISIONAL",
        "fecha_inicio": formatear_fecha(salida.fecha_salida),
        "fecha_fin": formatear_fecha(salida.fecha_fin_provisional),
        "finalidad_uso": "DE APOYO A LAS FUNCIONES SUSTANTIVAS",

        "num_oficio": "Solicitud de la DGAOPR",
        "nombre_persona": nombre_persona,
        "num_expediente": "JSG / DV",
        "num_licencia": valor(persona.num_licencia),
        "rfc": valor(persona.rfc),
        "area_adscripcion": "Dirección General de Apoyo al Ordenamiento y la Propiedad Rural",
        "cargo": valor(persona.cargo),
        "vigencia_licencia": valor(persona.vigencia_licencia),
        "tipo_licencia": valor(persona.tipo_licencia),

        "estado_llantas": valor(salida.estado_llantas_salida),
        "nivel_gasolina": valor(salida.nivel_gasolina_salida),
        "km_odometro_salida": valor(salida.km_odometro_salida),

        "num_tarjeta_gasolina": valor(vehiculo.num_tarjeta_gasolina, ""),

        "gasolina_0": check(normalizar_gasolina(salida.nivel_gasolina_salida), "0"),
        "gasolina_14": check(normalizar_gasolina(salida.nivel_gasolina_salida), "1/4"),
        "gasolina_12": check(normalizar_gasolina(salida.nivel_gasolina_salida), "1/2"),
        "gasolina_34": check(normalizar_gasolina(salida.nivel_gasolina_salida), "3/4"),
        "gasolina_44": check(normalizar_gasolina(salida.nivel_gasolina_salida), "4/4"),

        "llantas_14": check(normalizar_llantas(salida.estado_llantas_salida), "1/4"),
        "llantas_12": check(normalizar_llantas(salida.estado_llantas_salida), "1/2"),
        "llantas_34": check(normalizar_llantas(salida.estado_llantas_salida), "3/4"),
        "llantas_44": check(normalizar_llantas(salida.estado_llantas_salida), "4/4"),
    }

    contexto.update(generar_checks_condiciones(condiciones or []))
    contexto.update(generar_checks_inventario(inventario or []))

    if os.path.exists(ruta_template):
        documento = DocxTemplate(ruta_template)
        documento.render(contexto)
        documento.save(ruta_archivo)
    else:
        _generar_resguardo_basico(ruta_archivo, contexto, condiciones or [], inventario or [])

    return nombre_archivo, os.path.abspath(ruta_archivo)
